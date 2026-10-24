"""
Product/Inventory Management View & Add/Edit Product Dialog.
"""
import os
import csv
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QLineEdit,
    QTableWidget, QTableWidgetItem, QHeaderView, QFrame, QComboBox,
    QDialog, QFormLayout, QDoubleSpinBox, QSpinBox, QTextEdit,
    QFileDialog
)
from PySide6.QtCore import Qt, Signal, QTimer
from PySide6.QtGui import QFont

from app.services.billing import product_service
from app.ui.theme import theme_manager
from app.ui.theme.widgets import (
    PageHeaderWidget, CardFrame, StatusBadgeLabel, ThemedMessageBox,
    configure_dialog_window_flags
)


class ProductDialog(QDialog):
    """Add/Edit product dialog."""
    def __init__(self, product_data=None, parent=None):
        super().__init__(parent)
        self.product_data = product_data
        self.setWindowTitle("Edit Product" if product_data else "Add New Product")
        configure_dialog_window_flags(self, min_width=520, min_height=560)
        self.init_ui()
        if product_data:
            self.populate(product_data)

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(16)

        title = QLabel("Edit Product" if self.product_data else "New Product")
        title.setFont(theme_manager.get_font("title", bold=True))
        layout.addWidget(title)

        form = QFormLayout()
        form.setSpacing(12)

        self.inp_name = QLineEdit()
        self.inp_name.setMaxLength(200)
        self.inp_name.setPlaceholderText("Product name *")
        form.addRow("Name *", self.inp_name)

        self.inp_hsn = QLineEdit()
        self.inp_hsn.setMaxLength(20)
        self.inp_hsn.setPlaceholderText("e.g. 8504")
        form.addRow("HSN Code", self.inp_hsn)

        self.inp_category = QLineEdit()
        self.inp_category.setMaxLength(100)
        self.inp_category.setPlaceholderText("e.g. Electrical, Hardware")
        form.addRow("Category", self.inp_category)

        self.cmb_unit = QComboBox()
        self.cmb_unit.addItems(["PCS", "KGS", "BOX", "MTR", "NOS", "LTR", "SET", "PKT"])
        form.addRow("Unit", self.cmb_unit)

        self.spn_price = QDoubleSpinBox()
        self.spn_price.setRange(0, 99999999.99)
        self.spn_price.setDecimals(2)
        self.spn_price.setPrefix("₹ ")
        form.addRow("Selling Price", self.spn_price)

        self.spn_purchase = QDoubleSpinBox()
        self.spn_purchase.setRange(0, 99999999.99)
        self.spn_purchase.setDecimals(2)
        self.spn_purchase.setPrefix("₹ ")
        form.addRow("Purchase Price", self.spn_purchase)

        self.cmb_tax = QComboBox()
        self.cmb_tax.addItems(["0", "5", "12", "18", "28"])
        self.cmb_tax.setCurrentText("18")
        form.addRow("GST Rate %", self.cmb_tax)

        self.spn_stock = QDoubleSpinBox()
        self.spn_stock.setRange(0, 9999999)
        self.spn_stock.setDecimals(3)
        form.addRow("Stock Qty", self.spn_stock)

        self.spn_threshold = QDoubleSpinBox()
        self.spn_threshold.setRange(0, 9999999)
        self.spn_threshold.setDecimals(3)
        self.spn_threshold.setValue(10)
        form.addRow("Low Stock Threshold", self.spn_threshold)

        self.inp_barcode = QLineEdit()
        self.inp_barcode.setMaxLength(100)
        self.inp_barcode.setPlaceholderText("Optional barcode")
        form.addRow("Barcode", self.inp_barcode)

        self.inp_desc = QTextEdit()
        self.inp_desc.setMaximumHeight(60)
        self.inp_desc.setPlaceholderText("Optional description")
        form.addRow("Description", self.inp_desc)

        layout.addLayout(form)

        # Buttons
        btn_layout = QHBoxLayout()
        btn_layout.addStretch()

        btn_cancel = QPushButton("Cancel")
        btn_cancel.setProperty("variant", "secondary")
        btn_cancel.setFixedHeight(36)
        btn_cancel.clicked.connect(self.reject)
        btn_layout.addWidget(btn_cancel)

        btn_save = QPushButton("Save Product")
        btn_save.setIcon(theme_manager.get_icon("save", token_color="text_on_accent", size=16))
        btn_save.setProperty("variant", "primary")
        btn_save.setFixedHeight(36)
        btn_save.clicked.connect(self.accept)
        btn_layout.addWidget(btn_save)

        layout.addLayout(btn_layout)

    def populate(self, data):
        self.inp_name.setText(data.get("name", ""))
        self.inp_hsn.setText(data.get("hsn_code", "") or "")
        self.inp_category.setText(data.get("category", "") or "")
        idx = self.cmb_unit.findText(data.get("unit", "PCS"))
        if idx >= 0:
            self.cmb_unit.setCurrentIndex(idx)
        self.spn_price.setValue(data.get("price", 0))
        self.spn_purchase.setValue(data.get("purchase_price", 0) or 0)
        self.cmb_tax.setCurrentText(str(int(data.get("tax_rate", 18))))
        self.spn_stock.setValue(data.get("stock_quantity", 0))
        self.spn_threshold.setValue(data.get("low_stock_threshold", 10))
        self.inp_barcode.setText(data.get("barcode", "") or "")
        self.inp_desc.setPlainText(data.get("description", "") or "")

    def get_data(self):
        return {
            "name": self.inp_name.text().strip(),
            "hsn_code": self.inp_hsn.text().strip(),
            "category": self.inp_category.text().strip(),
            "unit": self.cmb_unit.currentText(),
            "price": self.spn_price.value(),
            "purchase_price": self.spn_purchase.value(),
            "tax_rate": float(self.cmb_tax.currentText()),
            "stock_quantity": self.spn_stock.value(),
            "low_stock_threshold": self.spn_threshold.value(),
            "barcode": self.inp_barcode.text().strip(),
            "description": self.inp_desc.toPlainText().strip(),
        }


class ProductView(QWidget):
    """Product/Inventory management page."""
    navigate_to = Signal(str, dict)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(16)

        # Header
        self.header_widget = PageHeaderWidget(
            title="Products & Inventory",
            subtitle="Manage product catalog, pricing, stock levels, and GST rates",
            action_text="Add Product",
            action_icon="plus",
            action_callback=self.open_add_dialog
        )

        btn_edit = QPushButton("Edit Selected")
        btn_edit.setIcon(theme_manager.get_icon("edit", size=16))
        btn_edit.setProperty("variant", "secondary")
        btn_edit.setFixedHeight(36)
        btn_edit.clicked.connect(self.open_edit_dialog)
        self.header_widget.add_custom_action(btn_edit)

        btn_export_csv = QPushButton("Export CSV")
        btn_export_csv.setIcon(theme_manager.get_icon("download", size=16))
        btn_export_csv.setProperty("variant", "secondary")
        btn_export_csv.setFixedHeight(36)
        btn_export_csv.clicked.connect(self.export_csv)
        self.header_widget.add_custom_action(btn_export_csv)

        btn_export_excel = QPushButton("Export Excel")
        btn_export_excel.setIcon(theme_manager.get_icon("download", size=16))
        btn_export_excel.setProperty("variant", "secondary")
        btn_export_excel.setFixedHeight(36)
        btn_export_excel.clicked.connect(self.export_excel)
        self.header_widget.add_custom_action(btn_export_excel)

        btn_refresh = QPushButton("Refresh")
        btn_refresh.setIcon(theme_manager.get_icon("refresh", size=16))
        btn_refresh.setProperty("variant", "secondary")
        btn_refresh.setFixedHeight(36)
        btn_refresh.clicked.connect(self.load_products)
        self.header_widget.add_custom_action(btn_refresh)

        layout.addWidget(self.header_widget)

        # Filters
        filter_frame = CardFrame()
        filter_layout = QHBoxLayout(filter_frame)
        filter_layout.setContentsMargins(12, 8, 12, 8)
        filter_layout.setSpacing(12)

        self.inp_search = QLineEdit()
        self.inp_search.setPlaceholderText("Search products by name, HSN, barcode...")
        self.search_timer = QTimer(self)
        self.search_timer.setSingleShot(True)
        self.search_timer.timeout.connect(self.load_products)
        self.inp_search.textChanged.connect(lambda: self.search_timer.start(200))
        filter_layout.addWidget(self.inp_search, 2)

        self.cmb_filter = QComboBox()
        self.cmb_filter.addItems(["All Products", "Low Stock Only"])
        self.cmb_filter.currentTextChanged.connect(self.load_products)
        filter_layout.addWidget(self.cmb_filter)

        layout.addWidget(filter_frame)

        # Products Table
        self.table = QTableWidget()
        self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels([
            "Name", "HSN", "Category", "Unit", "Price (₹)", "Purchase (₹)",
            "GST %", "Stock", "Status"
        ])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        for i in range(1, 9):
            self.table.horizontalHeader().setSectionResizeMode(i, QHeaderView.Interactive)
        self.table.setColumnWidth(1, 90)   # HSN
        self.table.setColumnWidth(2, 120)  # Category
        self.table.setColumnWidth(3, 75)   # Unit
        self.table.setColumnWidth(4, 110)  # Price
        self.table.setColumnWidth(5, 110)  # Purchase
        self.table.setColumnWidth(6, 75)   # GST %
        self.table.setColumnWidth(7, 95)   # Stock
        self.table.setColumnWidth(8, 100)  # Status

        self.table.verticalHeader().setDefaultSectionSize(36)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.doubleClicked.connect(self.open_edit_dialog)
        layout.addWidget(self.table, 1)

        # Bottom action bar
        bottom = QHBoxLayout()
        self.lbl_count = QLabel("0 products")
        self.lbl_count.setFont(theme_manager.get_font("caption"))
        bottom.addWidget(self.lbl_count)
        bottom.addStretch()

        btn_delete = QPushButton("Delete Selected")
        btn_delete.setIcon(theme_manager.get_icon("trash", token_color="danger", size=16))
        btn_delete.setProperty("variant", "danger")
        btn_delete.clicked.connect(self.delete_selected)
        bottom.addWidget(btn_delete)
        layout.addLayout(bottom)

        self.load_products()

    def load_products(self):
        search = self.inp_search.text().strip() if hasattr(self, 'inp_search') else None
        low_stock = self.cmb_filter.currentText() == "Low Stock Only" if hasattr(self, 'cmb_filter') else False

        products = product_service.get_all_products(search=search, low_stock_only=low_stock)
        self.table.setUpdatesEnabled(False)
        try:
            self.table.setRowCount(len(products))
            for row, p in enumerate(products):
                self.table.setItem(row, 0, QTableWidgetItem(p["name"]))
                self.table.setItem(row, 1, QTableWidgetItem(p.get("hsn_code") or "—"))
                self.table.setItem(row, 2, QTableWidgetItem(p.get("category") or "—"))
                self.table.setItem(row, 3, QTableWidgetItem(p.get("unit", "PCS")))

                it_price = QTableWidgetItem(f"₹{p['price']:,.2f}")
                it_price.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.table.setItem(row, 4, it_price)

                pp = p.get("purchase_price") or 0
                it_pp = QTableWidgetItem(f"₹{pp:,.2f}")
                it_pp.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.table.setItem(row, 5, it_pp)

                it_gst = QTableWidgetItem(f"{p['tax_rate']:.0f}%")
                it_gst.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(row, 6, it_gst)

                stock = p.get("stock_quantity", 0)
                threshold = p.get("low_stock_threshold", 10)
                it_stock = QTableWidgetItem(f"{stock:,.1f}")
                it_stock.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.table.setItem(row, 7, it_stock)

                if stock <= threshold:
                    badge = StatusBadgeLabel("LOW STOCK", status="danger")
                else:
                    badge = StatusBadgeLabel("IN STOCK", status="success")
                self.table.setCellWidget(row, 8, badge)

                self.table.item(row, 0).setData(Qt.UserRole, p["id"])
        finally:
            self.table.setUpdatesEnabled(True)

        self.lbl_count.setText(f"{len(products)} products")

    def open_add_dialog(self):
        dlg = ProductDialog(parent=self)
        if dlg.exec() == QDialog.Accepted:
            data = dlg.get_data()
            if not data["name"]:
                ThemedMessageBox.warning(self, "Validation", "Product name is required.")
                return
            try:
                product_service.create_product(data)
                self.load_products()
            except Exception as e:
                ThemedMessageBox.critical(self, "Error", str(e))

    def open_edit_dialog(self):
        row = self.table.currentRow()
        if row < 0:
            return
        p_id = self.table.item(row, 0).data(Qt.UserRole)
        p = product_service.get_product_by_id(p_id)
        if not p:
            return

        dlg = ProductDialog(product_data=p, parent=self)
        if dlg.exec() == QDialog.Accepted:
            data = dlg.get_data()
            try:
                product_service.update_product(p_id, data)
                self.load_products()
            except Exception as e:
                ThemedMessageBox.critical(self, "Error", str(e))

    def delete_selected(self):
        row = self.table.currentRow()
        if row < 0:
            ThemedMessageBox.information(self, "Info", "Select a product to delete.")
            return

        p_id = self.table.item(row, 0).data(Qt.UserRole)
        name = self.table.item(row, 0).text()

        if ThemedMessageBox.question(
            self,
            "Confirm Delete",
            f"Are you sure you want to delete product '{name}'?"
        ):
            product_service.delete_product(p_id)
            self.load_products()

    def export_csv(self):
        products = product_service.get_all_products()
        if not products:
            ThemedMessageBox.information(self, "No Data", "No products to export.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Products to CSV",
            os.path.expanduser("~/Documents/products_export.csv"),
            "CSV Files (*.csv)"
        )
        if not file_path:
            return

        try:
            fieldnames = [
                "ID", "Name", "HSN Code", "Category", "Unit", "Selling Price",
                "Purchase Price", "GST Rate (%)", "Stock Quantity",
                "Low Stock Threshold", "Barcode", "Description"
            ]

            with open(file_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(fieldnames)

                for p in products:
                    writer.writerow([
                        p.get("id", ""),
                        p.get("name", ""),
                        p.get("hsn_code", "") or "",
                        p.get("category", "") or "",
                        p.get("unit", "PCS"),
                        f"{p.get('price', 0):.2f}",
                        f"{p.get('purchase_price', 0) or 0:.2f}",
                        f"{p.get('tax_rate', 18):.1f}",
                        f"{p.get('stock_quantity', 0):.2f}",
                        f"{p.get('low_stock_threshold', 10):.2f}",
                        p.get("barcode", "") or "",
                        p.get("description", "") or ""
                    ])

            ThemedMessageBox.information(
                self,
                "Success",
                f"Products exported successfully to:\n{file_path}"
            )
        except Exception as e:
            ThemedMessageBox.critical(self, "Error", f"Failed to export CSV: {str(e)}")

    def export_excel(self):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            ThemedMessageBox.warning(
                self,
                "Missing Dependency",
                "openpyxl library is required to export to Excel.\nPlease install it or use CSV export."
            )
            return

        products = product_service.get_all_products()
        if not products:
            ThemedMessageBox.information(self, "No Data", "No products to export.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Products to Excel",
            os.path.expanduser("~/Documents/products_export.xlsx"),
            "Excel Files (*.xlsx)"
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Products Catalog"
            ws.views.sheetView[0].showGridLines = True

            accent_hex = theme_manager.get_export_accent().lstrip("#")
            header_fill = PatternFill(start_color=accent_hex, end_color=accent_hex, fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

            headers = [
                "ID", "Product Name", "HSN Code", "Category", "Unit",
                "Selling Price (₹)", "Purchase Price (₹)", "GST %",
                "Current Stock", "Min Threshold", "Barcode", "Description"
            ]

            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.row_dimensions[1].height = 28

            from app.services.export_service import sanitize_excel_cell, setup_excel_print_and_freeze
            setup_excel_print_and_freeze(ws, freeze_cell="A2", title_rows="1:1", landscape=True)

            thin_border = Border(
                left=Side(style='thin', color='D4D4D8'),
                right=Side(style='thin', color='D4D4D8'),
                top=Side(style='thin', color='D4D4D8'),
                bottom=Side(style='thin', color='D4D4D8')
            )

            data_font = Font(name="Calibri", size=10)
            low_stock_font = Font(name="Calibri", size=10, bold=True, color="DC2626")

            for r_idx, p in enumerate(products, start=2):
                stock = float(p.get("stock_quantity", 0))
                thresh = float(p.get("low_stock_threshold", 10))
                is_low = stock <= thresh

                row_data = [
                    p.get("id", ""),
                    sanitize_excel_cell(p.get("name", "")),
                    sanitize_excel_cell(p.get("hsn_code", "") or "-"),
                    sanitize_excel_cell(p.get("category", "") or "-"),
                    sanitize_excel_cell(p.get("unit", "PCS")),
                    float(p.get("price", 0)),
                    float(p.get("purchase_price", 0) or 0),
                    float(p.get("tax_rate", 18)),
                    stock,
                    thresh,
                    sanitize_excel_cell(p.get("barcode", "") or "-"),
                    sanitize_excel_cell(p.get("description", "") or "")
                ]
                ws.append(row_data)
                ws.row_dimensions[r_idx].height = 20

                for c_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    cell.border = thin_border
                    cell.font = low_stock_font if (is_low and c_idx == 9) else data_font

                    if c_idx in [6, 7]:
                        cell.number_format = '₹#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif c_idx in [8, 9, 10]:
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif c_idx in [1, 5]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = str(cell.value or '')
                    if len(val) > max_len:
                        max_len = len(val)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

            wb.save(file_path)
            ThemedMessageBox.information(
                self,
                "Success",
                f"Products exported successfully to:\n{file_path}"
            )
        except Exception as e:
            ThemedMessageBox.critical(self, "Error", f"Failed to export Excel: {str(e)}")
