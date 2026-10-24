"""
GST Reports & Feeder Summaries View (GSTR-1 Outward Sales & GSTR-3B Inward ITC)
"""
from datetime import date
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTabWidget, QTableWidget, QTableWidgetItem, QHeaderView,
    QDateEdit, QFrame
)
from PySide6.QtCore import Qt, QDate, Signal
from PySide6.QtGui import QColor, QFont
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.config import EXPORTS_DIR
import os

from app.services.report_service import get_gst_summary
from app.ui.widgets.stat_card import StatCard
from app.ui.theme import theme_manager
from app.ui.theme.widgets import (
    PageHeaderWidget, CardFrame, ThemedMessageBox
)
from app.ui.async_worker import run_in_background
from app.ui.save_dialog import execute_export_pipeline


class GSTTabWidget(QWidget):
    def __init__(self, report_type="Sales", parent=None):
        super().__init__(parent)
        self.report_type = report_type
        self.current_data = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 12, 8, 8)
        layout.setSpacing(12)

        # KPI Row
        kpi_layout = QHBoxLayout()
        kpi_layout.setSpacing(10)
        self.card_taxable = StatCard("TAXABLE TURNOVER", "₹0.00", "Net value before tax", token_name="accent")
        self.card_cgst = StatCard("TOTAL CGST", "₹0.00", "Central tax", token_name="success")
        self.card_sgst = StatCard("TOTAL SGST", "₹0.00", "State tax", token_name="success")
        self.card_igst = StatCard("TOTAL IGST", "₹0.00", "Inter-state tax", token_name="info")
        self.card_total = StatCard("TOTAL INVOICE VALUE", "₹0.00", "Gross value", token_name="text_primary")

        kpi_layout.addWidget(self.card_taxable)
        kpi_layout.addWidget(self.card_cgst)
        kpi_layout.addWidget(self.card_sgst)
        kpi_layout.addWidget(self.card_igst)
        kpi_layout.addWidget(self.card_total)
        layout.addLayout(kpi_layout)

        # Filter
        f_card = CardFrame()
        f_layout = QHBoxLayout(f_card)
        f_layout.setContentsMargins(12, 8, 12, 8)
        f_layout.setSpacing(10)

        lbl_from = QLabel("From:")
        lbl_from.setFont(theme_manager.get_font("label", bold=True))
        f_layout.addWidget(lbl_from)
        self.dt_from = QDateEdit()
        self.dt_from.setCalendarPopup(True)
        self.dt_from.setDate(QDate(date.today().year, date.today().month, 1))
        self.dt_from.setDisplayFormat("dd-MM-yyyy")
        f_layout.addWidget(self.dt_from)

        lbl_to = QLabel("To:")
        lbl_to.setFont(theme_manager.get_font("label", bold=True))
        f_layout.addWidget(lbl_to)
        self.dt_to = QDateEdit()
        self.dt_to.setCalendarPopup(True)
        self.dt_to.setDate(QDate.currentDate())
        self.dt_to.setDisplayFormat("dd-MM-yyyy")
        f_layout.addWidget(self.dt_to)

        from app.ui.widgets.quick_date_selector import QuickDateSelector
        self.quick_dates = QuickDateSelector(self.dt_from, self.dt_to, on_change=self.load_data)
        f_layout.addWidget(self.quick_dates)

        btn_filter = QPushButton("Load Summary")
        btn_filter.setIcon(theme_manager.get_icon("refresh", size=16))
        btn_filter.setProperty("variant", "primary")
        btn_filter.setFixedHeight(34)
        btn_filter.clicked.connect(self.load_data)
        f_layout.addWidget(btn_filter)

        f_layout.addStretch()

        if self.report_type == "Sales":
            self.btn_json = QPushButton("Export GSTR-1 JSON")
            self.btn_json.setIcon(theme_manager.get_icon("download", size=16))
            self.btn_json.setProperty("variant", "secondary")
            self.btn_json.setFixedHeight(34)
            self.btn_json.setToolTip("Export GSTR-1 JSON compatible with official GST Offline Tool")
            self.btn_json.clicked.connect(self.export_json)
            f_layout.addWidget(self.btn_json)

        self.btn_excel = QPushButton("Export for CA (Excel)")
        self.btn_excel.setIcon(theme_manager.get_icon("download", size=16))
        self.btn_excel.setProperty("variant", "secondary")
        self.btn_excel.setFixedHeight(34)
        self.btn_excel.clicked.connect(self.export_excel)
        f_layout.addWidget(self.btn_excel)

        layout.addWidget(f_card)

        # Table
        self.table = QTableWidget()
        self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels([
            "Date", "Invoice / Bill No.", "Party Name", "GSTIN",
            "Taxable (₹)", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Total (₹)"
        ])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        for i in range(3, 9):
            self.table.horizontalHeader().setSectionResizeMode(i, QHeaderView.ResizeToContents)

        self.table.verticalHeader().setDefaultSectionSize(36)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        layout.addWidget(self.table, 1)

        self.load_data()

    def load_data(self):
        q_f = self.dt_from.date()
        q_t = self.dt_to.date()
        from_d = date(q_f.year(), q_f.month(), q_f.day())
        to_d = date(q_t.year(), q_t.month(), q_t.day())

        data = get_gst_summary(from_d, to_d, report_type=self.report_type)
        self.current_data = data

        self.card_taxable.set_value(f"₹{data['total_taxable']:,.2f}")
        self.card_cgst.set_value(f"₹{data['total_cgst']:,.2f}")
        self.card_sgst.set_value(f"₹{data['total_sgst']:,.2f}")
        self.card_igst.set_value(f"₹{data['total_igst']:,.2f}")
        self.card_total.set_value(f"₹{data['total_value']:,.2f}", f"{data['count']} Invoices")

        rows = data["rows"]
        total_rows = len(rows) + 1
        self.table.setUpdatesEnabled(False)
        try:
            self.table.setRowCount(total_rows)

            for r_idx, r in enumerate(rows):
                d_str = r["voucher_date"].strftime("%d-%m-%Y") if hasattr(r["voucher_date"], "strftime") else str(r["voucher_date"])

                it_d = QTableWidgetItem(d_str)
                it_v = QTableWidgetItem(r["voucher_number"])
                it_p = QTableWidgetItem(r["party_name"])
                it_g = QTableWidgetItem(r["party_gstin"] or "—")

                self.table.setItem(r_idx, 0, it_d)
                self.table.setItem(r_idx, 1, it_v)
                self.table.setItem(r_idx, 2, it_p)
                self.table.setItem(r_idx, 3, it_g)

                for c_idx, key in enumerate(["taxable_value", "cgst", "sgst", "igst", "total_value"], start=4):
                    val = r[key]
                    it = QTableWidgetItem(f"₹{val:,.2f}" if val > 0 else "—")
                    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                    self.table.setItem(r_idx, c_idx, it)

            # Grand Total Row
            tot_idx = total_rows - 1
            it_t_lbl = QTableWidgetItem("GRAND TOTALS")
            it_t_lbl.setFont(theme_manager.get_font("body", bold=True))
            self.table.setItem(tot_idx, 2, it_t_lbl)

            for c_idx, key in enumerate(["total_taxable", "total_cgst", "total_sgst", "total_igst", "total_value"], start=4):
                val = data[key]
                it = QTableWidgetItem(f"₹{val:,.2f}")
                it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                it.setFont(theme_manager.get_font("body", bold=True))
                if key == "total_value":
                    it.setForeground(theme_manager.qcolor("accent"))
                self.table.setItem(tot_idx, c_idx, it)
        finally:
            self.table.setUpdatesEnabled(True)

    def export_excel(self):
        if not self.current_data or not self.current_data.get("rows"):
            ThemedMessageBox.warning(self, "No Data", "No GST records to export.")
            return

        cur_data = self.current_data
        rep_type = self.report_type
        suggested = f"GST_{rep_type}_{cur_data['from_date']}_to_{cur_data['to_date']}"

        def _generate(out_p):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"GST {rep_type}"

            # Write Header
            ws.cell(row=1, column=1, value=f"GST {rep_type.upper()} SUMMARY REPORT").font = Font(bold=True, size=14)
            ws.cell(row=2, column=1, value=f"Period: {cur_data['from_date']} to {cur_data['to_date']}").font = Font(italic=True)

            accent_hex = theme_manager.get_export_accent().lstrip("#")
            header_fill = PatternFill(start_color=accent_hex, end_color=accent_hex, fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            headers = ["Date", "Invoice / Bill No.", "Party Name", "GSTIN", "Taxable Value (₹)", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Total Value (₹)"]
            for col_num, h in enumerate(headers, start=1):
                c = ws.cell(row=4, column=col_num, value=h)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="right" if "₹" in h else "left")

            from app.services.export_service import sanitize_excel_cell, setup_excel_print_and_freeze, INR_EXCEL_FORMAT
            setup_excel_print_and_freeze(ws, freeze_cell="A5", title_rows="4:4", landscape=True)

            row_num = 4
            for r in cur_data["rows"]:
                row_num += 1
                d_str = r["voucher_date"].strftime("%d-%m-%Y") if hasattr(r["voucher_date"], "strftime") else str(r["voucher_date"])
                ws.cell(row=row_num, column=1, value=d_str)
                ws.cell(row=row_num, column=2, value=sanitize_excel_cell(r["voucher_number"]))
                ws.cell(row=row_num, column=3, value=sanitize_excel_cell(r["party_name"]))
                ws.cell(row=row_num, column=4, value=sanitize_excel_cell(r["party_gstin"]))

                for idx, k in enumerate(["taxable_value", "cgst", "sgst", "igst", "total_value"], start=5):
                    c = ws.cell(row=row_num, column=idx, value=r[k])
                    c.number_format = INR_EXCEL_FORMAT

            # Total row
            row_num += 1
            ws.cell(row=row_num, column=3, value="GRAND TOTALS").font = Font(bold=True)
            for idx, k in enumerate(["total_taxable", "total_cgst", "total_sgst", "total_igst", "total_value"], start=5):
                c = ws.cell(row=row_num, column=idx, value=cur_data[k])
                c.font = Font(bold=True)
                c.number_format = INR_EXCEL_FORMAT

            from app.services.export_service import autofit_excel_columns
            autofit_excel_columns(ws)

            wb.save(str(out_p))
            return out_p

        execute_export_pipeline(
            parent=self,
            suggested_name=suggested,
            kind="xlsx",
            generator_fn=_generate,
            trigger_button=self.btn_excel
        )

    def export_json(self):
        if not self.current_data or self.report_type != "Sales":
            return

        q_f = self.dt_from.date()
        q_t = self.dt_to.date()
        from_d = date(q_f.year(), q_f.month(), q_f.day())
        to_d = date(q_t.year(), q_t.month(), q_t.day())
        suggested = f"GSTR1_{from_d.strftime('%Y%m%d')}_to_{to_d.strftime('%Y%m%d')}"

        def _generate(out_p):
            from app.services.gst_report_service import export_gstr1_offline_json
            return export_gstr1_offline_json(from_d, to_d, output_path=str(out_p))

        execute_export_pipeline(
            parent=self,
            suggested_name=suggested,
            kind="json",
            generator_fn=_generate,
            trigger_button=self.btn_json if hasattr(self, "btn_json") else None
        )


class GSTReportView(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(14)

        self.header_widget = PageHeaderWidget(
            title="GST Returns & Tax Summaries",
            subtitle="GSTR-1 outward supplies & GSTR-3B inward input tax credit (ITC) feeder reports"
        )
        layout.addWidget(self.header_widget)

        self.tabs = QTabWidget()
        self.tabs.addTab(GSTTabWidget("Sales", parent=self), "Outward Supplies (GSTR-1 Sales)")
        self.tabs.addTab(GSTTabWidget("Purchases", parent=self), "Inward Supplies (GSTR-3B Purchase ITC)")
        layout.addWidget(self.tabs, 1)
