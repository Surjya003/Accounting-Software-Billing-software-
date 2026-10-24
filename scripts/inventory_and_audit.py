"""
BillFlow Suite - Refined Inventory and Cleanup Candidate Analysis
Strictly adheres to Hard Rule 2 (never touch DBs, user data, core code, tests, docs).
Produces cleanup_candidates.xlsx with dual evidence for every DELETE candidate.
"""
import os
import sys
import re
import hashlib
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Set, Tuple, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(".").resolve()

PROTECTED_EXACT_FILES = {
    "build_exe.spec",
    "installer.iss",
    "requirements.txt",
    "requirements_dev.txt",
    "pyproject.toml",
    "setup.py",
    "run_app.bat",
    "README.md",
    "USER_GUIDE.md",
    "ADMIN_GUIDE.md",
    "COMPLIANCE.md",
    "KNOWN_LIMITATIONS.md",
    "CHANGELOG.md",
    "THEMING.md",
    "UI_GUIDELINES.md",
    ".gitignore",
    ".env",
    "COMMERCIAL_GUIDE.md",
    "AUDIT_REPORT.md",
}

CACHE_JUNK_PATTERNS = [
    r"__pycache__",
    r"\.pyc$",
    r"\.pytest_cache",
    r"\.mypy_cache",
    r"\.ruff_cache",
    r"\.coverage",
    r"htmlcov",
    r"^build[/\\]",
    r"^dist[/\\]",
    r"\.egg-info",
    r"\.idea",
    r"Thumbs\.db$",
    r"\.DS_Store$",
    r"\.log$",
    r"\.tmp$",
    r"\.bak$",
    r"~$",
    r"\.orig$",
    r"\.rej$",
    r"\.swp$",
]

def sha256_file(filepath: Path) -> str:
    h = hashlib.sha256()
    try:
        with open(filepath, "rb") as f:
            while chunk := f.read(65536):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return ""

def scan_repository():
    file_entries = []
    for root, dirs, files in os.walk(ROOT):
        rel_root = Path(root).relative_to(ROOT)
        is_git = rel_root.parts and rel_root.parts[0] == ".git"

        for f in files:
            p = Path(root) / f
            rel = p.relative_to(ROOT)
            rel_str = str(rel).replace("\\", "/")
            try:
                st = p.stat()
                sz = st.st_size
                mtime = datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            except Exception:
                sz = 0
                mtime = "unknown"

            file_entries.append({
                "path": p,
                "rel": rel_str,
                "size": sz,
                "mtime": mtime,
                "is_git": is_git,
            })
    return file_entries

def classify_entry(entry: Dict[str, Any]):
    rel = entry["rel"]
    sz = entry["size"]

    # 1. Git internals
    if entry["is_git"]:
        return "F", "Protected (.git internal)", "Git version control tree", "KEEP", "Low"

    # 2. Database & User Data (Rule 2)
    if re.search(r"\.(db|sqlite\d*|db-wal|db-shm|enc)$", rel, re.IGNORECASE) or rel.startswith("data/") or rel.startswith("backups/"):
        return "F", "Protected (Database / User Data / backups)", "Rule 2: Never modify database or user data", "KEEP", "Low"

    # 3. Core App Code, Assets, Themes, Tests (Rule 2)
    if rel.startswith("app/") or rel.startswith("assets/") or rel.startswith("themes/") or rel.startswith("tests/"):
        if "__pycache__" in rel or rel.endswith(".pyc"):
            return "A", "Python bytecode cache file", "Ephemeral compiled cache; regenerated automatically", "DELETE", "Low"
        return "F", "Protected (Core app source / assets / tests)", "Rule 2: Active application code and verification suite", "KEEP", "Low"

    # 4. Release & QA Scripts
    if rel.startswith("scripts/"):
        if "__pycache__" in rel or rel.endswith(".pyc"):
            return "A", "Python bytecode cache file", "Ephemeral compiled cache; regenerated automatically", "DELETE", "Low"
        return "F", "Protected (Release/QA script)", "Rule 2: Verification and build tooling", "KEEP", "Low"

    # 5. Core Documentation & Packaging Configurations
    if rel in PROTECTED_EXACT_FILES:
        return "F", "Protected (Current documentation / packaging)", "Rule 2: Official project documentation and build specs", "KEEP", "Low"

    # 6. QA Reports & Evaluation Excel spreadsheets (Rule 2 explicitly protects QA reports)
    if rel.startswith("QA_") and rel.endswith(".xlsx"):
        return "E", "QA evaluation report", "Rule 2: Official QA delivery report", "KEEP", "Low"

    # 7. Submodule repository Billing-software
    if rel.startswith("Billing-software/"):
        return "G", "Separate git submodule repository", "Legacy repository tree; requires owner directive", "NEEDS OWNER", "High"

    # 8. Audit Evidence Folders explicitly protected in Rule 2
    if rel.startswith("docs/ui_audit/") or rel.startswith("docs/pdf_samples/") or rel.startswith("docs/theme_screenshots/") or rel.startswith("screenshots/full_audit/"):
        return "E", "UI Audit / PDF sample verification records", "Rule 2: Protected QA reports and visual proof artifacts", "KEEP", "Low"

    # 9. Stale full copy of app before theme
    if rel.startswith("backup_before_theme_app/"):
        return "C", "Duplicate full copy of app/ before theme refactor", "Stale duplicate; live source resides in app/", "DELETE", "Low"

    # 10. Cache Junk (Bucket A)
    for pat in CACHE_JUNK_PATTERNS:
        if re.search(pat, rel, re.IGNORECASE):
            return "A", f"Matches cache pattern: {pat}", "Build artifact / generated junk; safely rebuildable", "DELETE", "Low"

    # 11. Stale duplicate root files
    if rel == "redme.md":
        return "C", "Spelling duplicate of README.md", "Redundant lower-case typo duplicate; README.md is maintained", "DELETE", "Low"
    if rel == "testprompt.txt":
        return "B", "Scratch test prompt note", "Ad-hoc temporary scratch prompt file", "DELETE", "Low"
    if rel == "sample_proforma_invoice.pdf":
        return "B", "Stale sample PDF in root", "Temporary manual export artifact; live templates in app/services/pdf", "DELETE", "Low"

    # 12. Scratch directory
    if rel.startswith("scratch/"):
        return "B", "Scratch execution test artifact", "Temporary test output; safely regenerable", "DELETE", "Low"

    # 13. Exports directory
    if rel.startswith("exports/"):
        return "B", "Generated document export artifacts", "Generated test exports; live generation tested via suite", "DELETE", "Low"

    return "G", "Unclassified file", "Requires owner review", "NEEDS OWNER", "Medium"

def main():
    entries = scan_repository()
    classified = []
    bucket_counts = {"A": 0, "B": 0, "C": 0, "D": 0, "E": 0, "F": 0, "G": 0}
    bucket_sizes = {"A": 0, "B": 0, "C": 0, "D": 0, "E": 0, "F": 0, "G": 0}

    for e in entries:
        b, ev1, ev2, dec, risk = classify_entry(e)
        e["bucket"] = b
        e["ev1"] = ev1
        e["ev2"] = ev2
        e["decision"] = dec
        e["risk"] = risk
        bucket_counts[b] += 1
        bucket_sizes[b] += e["size"]
        classified.append(e)

    # Generate cleanup_candidates.xlsx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cleanup Candidates"

    headers = ["Path", "Size (Bytes)", "Size (KB)", "Bucket", "Evidence 1", "Evidence 2", "Last Modified", "Decision", "Risk"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    candidates = [e for e in classified if not e["is_git"]]
    candidates.sort(key=lambda x: (0 if x["decision"] == "DELETE" else 1, -x["size"]))

    for e in candidates:
        ws.append([
            e["rel"],
            e["size"],
            round(e["size"] / 1024.0, 2),
            e["bucket"],
            e["ev1"],
            e["ev2"],
            e["mtime"],
            e["decision"],
            e["risk"],
        ])

    wb.save("cleanup_candidates.xlsx")
    print(f"Updated cleanup_candidates.xlsx ({len(candidates)} non-git files cataloged).")

if __name__ == "__main__":
    main()
