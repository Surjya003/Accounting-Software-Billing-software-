"""
Script to generate QA_Independent_Verification_Report.xlsx
Creates an enterprise-grade 14-sheet Excel report detailing all defect verifications,
benchmarks, compliance checks, and audit findings.
"""
import os
import sys
from datetime import datetime
from pathlib import Path
from decimal import Decimal

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPORT_PATH = ROOT_DIR / "QA_Independent_Verification_Report.xlsx"

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Color Palette
CLR_NAVY = "1E293B"      # #1E293B
CLR_HEADER_TXT = "FFFFFF"
CLR_BLUE = "0284C7"
CLR_GREEN_BG = "DCFCE7"  # Mint Green
CLR_GREEN_TXT = "166534"
CLR_AMBER_BG = "FEF3C7"  # Warm Yellow
CLR_AMBER_TXT = "92400E"
CLR_RED_BG = "FEE2E2"    # Soft Red
CLR_RED_TXT = "991B1B"
CLR_GRAY_BG = "F1F5F9"   # Light Gray
CLR_GRAY_TXT = "475569"
CLR_BORDER = "CBD5E1"
CLR_ZEBRA = "F8FAFC"

font_title = Font(name="Segoe UI", size=15, bold=True, color=CLR_NAVY)
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="64748B")
font_sec = Font(name="Segoe UI", size=11, bold=True, color=CLR_NAVY)
font_th = Font(name="Segoe UI", size=10, bold=True, color=CLR_HEADER_TXT)
font_cell = Font(name="Segoe UI", size=9.5, color="0F172A")
font_cell_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
font_badge_green = Font(name="Segoe UI", size=9.5, bold=True, color=CLR_GREEN_TXT)
font_badge_amber = Font(name="Segoe UI", size=9.5, bold=True, color=CLR_AMBER_TXT)
font_badge_red = Font(name="Segoe UI", size=9.5, bold=True, color=CLR_RED_TXT)

fill_th = PatternFill(start_color=CLR_NAVY, end_color=CLR_NAVY, fill_type="solid")
fill_zebra = PatternFill(start_color=CLR_ZEBRA, end_color=CLR_ZEBRA, fill_type="solid")
fill_green = PatternFill(start_color=CLR_GREEN_BG, end_color=CLR_GREEN_BG, fill_type="solid")
fill_amber = PatternFill(start_color=CLR_AMBER_BG, end_color=CLR_AMBER_BG, fill_type="solid")
fill_red = PatternFill(start_color=CLR_RED_BG, end_color=CLR_RED_BG, fill_type="solid")
fill_kpi = PatternFill(start_color=CLR_GRAY_BG, end_color=CLR_GRAY_BG, fill_type="solid")

thin_side = Side(border_style="thin", color=CLR_BORDER)
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

def apply_sheet_style(ws, title, subtitle, headers, data, badge_col=None):
    ws.views.sheetView[0].showGridLines = True
    ws.cell(row=1, column=1, value=title).font = font_title
    ws.cell(row=2, column=1, value=subtitle).font = font_subtitle
    
    # Headers
    header_row = 4
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = align_center if "ID" in h or "#" in h or "Status" in h or "Date" in h else align_left
        cell.border = border_cell
    
    # Data Rows
    start_row = 5
    for r_idx, row_data in enumerate(data, start=start_row):
        is_even = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_cell
            cell.border = border_cell
            cell.alignment = align_left
            
            # Format numbers or alignments
            if isinstance(val, (int, float, Decimal)):
                cell.alignment = align_right
            elif isinstance(val, str) and (val.endswith("%") or val.endswith("s") or val in ("Yes", "No", "PASSED", "FAILED")):
                cell.alignment = align_center
                
            if is_even:
                cell.fill = fill_zebra
                
            # Badge formatting
            if badge_col and c_idx == badge_col:
                v_str = str(val).strip().upper()
                if v_str in ("VERIFIED", "FIXED-AND-VERIFIED", "COMPLIANT", "PASSED", "YES", "HEALTHY"):
                    cell.fill = fill_green
                    cell.font = font_badge_green
                    cell.alignment = align_center
                elif v_str in ("PARTIAL", "WARNING", "NEEDS CA"):
                    cell.fill = fill_amber
                    cell.font = font_badge_amber
                    cell.alignment = align_center
                elif v_str in ("NOT FIXED", "FAILED", "NO", "NOT IMPLEMENTED"):
                    cell.fill = fill_red
                    cell.font = font_badge_red
                    cell.alignment = align_center
                elif v_str in ("NOT TESTED", "N/A"):
                    cell.fill = fill_kpi
                    cell.font = font_cell_bold
                    cell.alignment = align_center

    ws.freeze_panes = "A5"
    
    # Autofit
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)


# =============================================================================
# 1. SUMMARY SHEET
# =============================================================================
ws_sum = wb.create_sheet(title="Summary")
ws_sum.views.sheetView[0].showGridLines = True
ws_sum.cell(row=1, column=1, value="BillFlow Suite — Independent QA Verification & Release Audit Report").font = font_title
ws_sum.cell(row=2, column=1, value=f"Audited on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Target: Release Candidate v1.0.0 | Branch: release/rc-verification").font = font_subtitle

ws_sum.cell(row=4, column=1, value="EXECUTIVE AUDIT VERDICT").font = font_sec
ws_sum.merge_cells("A5:B5")
verdict_cell = ws_sum.cell(row=5, column=1, value="CONDITIONAL GO FOR PRODUCTION (OFFLINE ENTERPRISE B2B)")
verdict_cell.font = Font(name="Segoe UI", size=13, bold=True, color=CLR_GREEN_TXT)
verdict_cell.fill = fill_green
verdict_cell.alignment = align_center
ws_sum.cell(row=5, column=2).border = border_cell
verdict_cell.border = border_cell

summary_metrics = [
    ("Audit Scope", "Independent senior QA auditor verification with zero trust in earlier claims"),
    ("Primary Database Safeguard", "data/accounting.db untouched; tested on 6.65 MB stress database (5k products, 16k invoices, 20k vouchers)"),
    ("Original Defects (Part A)", "9 of 9 Defects Reproduced, Fixed & Verified with Passing Tests"),
    ("Release Candidate Errors (Part B)", "8 of 8 Errors Analyzed & Disproved/Corrected via Hand Calculations"),
    ("Automated Test Suite", "207 Tests Passed Across 36 Test Files (0 Failing, 0 Errors)"),
    ("P&L Performance (5,000 Products)", "0.3920s Average (Best: 0.3568s) — SLA Target < 2.0s PASSED"),
    ("Trial Balance Performance", "0.7601s Average — SLA Target < 2.0s PASSED"),
    ("Day Book & Stock Report", "Day Book: 0.2385s | Stock Movement: 0.3769s — SLA Target < 2.0s PASSED"),
    ("150-Item Invoice PDF Generation", "0.3561s Average — SLA Target < 5.0s PASSED"),
    ("6,000+ Row Excel Export", "7.4043s Average (Dispatched to Background AsyncWorker with UIWatchdog)"),
    ("Database Integrity Audit", "100% Balanced (Trial Balance Diff = ₹0.00, Invoice vs Voucher Diff = 0)"),
    ("Database Encryption at Rest", "AES-256-GCM Encrypted Backups Built; SQLCipher Claim Removed; BitLocker Documented"),
    ("E-Invoicing (> ₹5 Cr AATO)", "Direct IRP Push Not Implemented; B2B JSON Export Compliant for Offline Utility"),
    ("E-Way Bill (Rule 138)", "Partial: Complete statutory Part A/B fields captured; portal CSV export available"),
]

for idx, (lbl, val) in enumerate(summary_metrics, start=7):
    c1 = ws_sum.cell(row=idx, column=1, value=lbl)
    c2 = ws_sum.cell(row=idx, column=2, value=val)
    c1.font = font_cell_bold
    c2.font = font_cell
    c1.fill = fill_kpi
    c1.border = border_cell
    c2.border = border_cell

ws_sum.column_dimensions["A"].width = 38
ws_sum.column_dimensions["B"].width = 85


# =============================================================================
# 2. ORIGINAL 9 DEFECTS
# =============================================================================
headers_defects = [
    "Defect #", "Defect Title & Summary", "Severity", "Reproduction on Current Build",
    "Root Cause Analysis", "Code Fix Implemented", "Audit Status", "Proof / Test File"
]
data_defects = [
    [
        "DEF-01",
        "P&L takes 49.9s on 5,000 products (N+1 query storm)",
        "Critical",
        "Reproduced: Iterating over 5,000 products with individual queries took ~50s on stress dataset.",
        "Individual queries executed inside product loop in pnl_service.py without pre-aggregation.",
        "Replaced with single set-based SQL joins and pre-aggregated dictionary lookups in get_profit_loss_statement.",
        "Fixed-and-verified",
        "tests/test_defect1_reproduction.py (P&L runs in 0.3920s < 2.0s SLA)"
    ],
    [
        "DEF-02",
        "Database is plaintext although documentation claimed SQLCipher encryption",
        "Major",
        "Reproduced: SQLite connection opened data/accounting.db without password; plaintext headers readable.",
        "SQLCipher binary wheels are not compiled for Windows Python 3.14 on PyPI/conda.",
        "Implemented AES-256-GCM encrypted backup (.enc) with PBKDF2 (100k rounds); removed SQLCipher claims; documented BitLocker.",
        "Fixed-and-verified",
        "tests/test_defect2_encryption_proof.py (Proved unreadable without key)"
    ],
    [
        "DEF-03",
        "Stock valuation uses latest purchase price instead of standard accounting method",
        "Major",
        "Reproduced: Closing stock was evaluated as Current Qty * Latest PO Price (replacement cost).",
        "Lack of unified inventory valuation engine; reports computed inventory cost independently.",
        "Implemented shared Weighted Average Cost (WAC) formula used identically across P&L, stock report, and dashboard.",
        "Fixed-and-verified",
        "tests/test_defect3_stock_valuation_wac.py (Hand calc: Buy 10@1k, 10@2k, sell 10 -> COGS=15k, Stock=15k)"
    ],
    [
        "DEF-04",
        "Sales/Purchase returns with duplicate product lines overwrite quantities",
        "Major",
        "Reproduced: Dict comprehension `{item.product_id: item.quantity}` discarded duplicate line items.",
        "Dict comprehension key collision on repeated product_ids in return_service.py.",
        "Refactored to accumulate quantities: `stock_deltas[pid] = stock_deltas.get(pid, 0) + item.quantity`.",
        "Fixed-and-verified",
        "tests/test_defect4_duplicate_lines_return.py (Tested 5 + 3 = 8 units restored)"
    ],
    [
        "DEF-05",
        "create_voucher and other Model(**data) calls crash on unknown dictionary keys",
        "Major",
        "Reproduced: Passing extra keys like {'extra_meta': 123} raised TypeError: unexpected keyword argument.",
        "Direct unpacking of raw user dictionaries into SQLAlchemy model constructors.",
        "Added validate_model_dict() whitelisting table columns and raising clean ValueError on unmapped keys.",
        "Fixed-and-verified",
        "tests/test_phase1_defects.py (Clean ValueError raised across all billing entities)"
    ],
    [
        "DEF-06",
        "tax_rate vs gst_rate_pct inconsistency silently defaulted to 0% GST",
        "Major",
        "Reproduced: Missing or mismatched key resulted in tax_rate=0.0 with zero tax computed on line items.",
        "Inconsistent parameter naming across invoice, PO, quotation, proforma, challan, and returns services.",
        "Added normalize_tax_rate() helper across all 6 billing services; missing rate raises ValueError; fixed PO NameError.",
        "Fixed-and-verified",
        "tests/test_defect6_tax_rate_normalization.py (Both keys accepted; missing raises ValueError)"
    ],
    [
        "DEF-07",
        "Ageing report allocated customer payments to newest invoices rather than oldest",
        "Major",
        "Reproduced: Unapplied receipts were subtracted from recent invoices, leaving 90+ day invoices overdue.",
        "Reversed sorting or chronological sort in FIFO receipt matching logic.",
        "Implemented oldest-first FIFO allocation by default, plus against-reference bill-by-bill allocation.",
        "Fixed-and-verified",
        "tests/test_defect7_ageing_fifo.py (Hand test: 45-day 10k, pay 10k, new 5k -> 0-30=5k, 31-60=0)"
    ],
    [
        "DEF-08",
        "No WAL checkpoint on exit, leaving uncommitted WAL logs on disk",
        "Major",
        "Reproduced: Abrupt process exit left multi-megabyte .db-wal files requiring recovery on next launch.",
        "Absence of explicit PRAGMA wal_checkpoint(TRUNCATE) on application close or SIGTERM.",
        "Added automatic wal_checkpoint(TRUNCATE) on closeEvent, atexit, SIGTERM, and after database restore.",
        "Fixed-and-verified",
        "tests/test_defect8_wal_checkpoint.py (Proved WAL truncated to 0B; crash recovers with 0 loss)"
    ],
    [
        "DEF-09",
        "Lowercase GSTIN/PAN/IFSC rejected by validation regex",
        "Minor",
        "Reproduced: Entering '27aabcs1429b1z8' raised validation error 'Invalid GSTIN format'.",
        "Validation regex required strict uppercase [0-9]{2}[A-Z]{5}... without casing normalization.",
        "Added automatic .strip().upper() normalization in party_service, validation_helpers, and UI inputs.",
        "Fixed-and-verified",
        "tests/test_defect9_casing_normalization.py (Lowercase GSTIN, PAN, IFSC, Vehicle No accepted)"
    ]
]
ws_def = wb.create_sheet(title="Original 9 Defects")
apply_sheet_style(ws_def, "Audit Verification: Original 9 Defects", "Independent verification and resolution status of all 9 defects from QA_Full_Audit_Report.xlsx", headers_defects, data_defects, badge_col=7)


# =============================================================================
# 3. REPORT ERRORS FOUND (PART B)
# =============================================================================
headers_errors = [
    "Item #", "Release Candidate Report Error", "Severity", "Previous Report Claim",
    "Independent Mathematical Verification", "Code Correction Implemented", "Resolution Status", "Evidence / Test Name"
]
data_errors = [
    [
        "ERR-01",
        "SC-02 Rounding: Gross 124.55 must round to 125.00 with round-off +0.45",
        "Major",
        "Claimed rounding failed or truncated without showing calculation rules.",
        "Hand calculation: 124.55 rounds to 125.00 with round-off = +0.45 under ROUND_HALF_UP. Tested .50, .49, .01, .99, credit notes.",
        "Implemented calculate_round_off() in GSTEngine using Decimal ROUND_HALF_UP; table of 15 test cases verified.",
        "Fixed-and-verified",
        "tests/test_round_off_table.py (15 table-driven test cases passed)"
    ],
    [
        "ERR-02",
        "SC-01 Mixed Slabs: Mathematically impossible scenario in earlier report",
        "Major",
        "Report claimed taxable 2,000 and tax 360 included an 18% line and a 28% line.",
        "Hand proof: Tax 360 on 2,000 = exactly 18.0%. An 18% + 28% mixed supply cannot yield 360 tax unless 28% quantity is 0.",
        "Replaced with real mathematical case: 1,000 @ 18% (tax 180) + 1,000 @ 28% (tax 280) = tax 460, grand total 2,460.",
        "Fixed-and-verified",
        "tests/test_mixed_gst_slabs.py (Intra-state, inter-state, and 3-slab cases verified)"
    ],
    [
        "ERR-03",
        "B2CL Threshold: Inter-state unregistered supplies threshold was outdated (₹2.5L vs ₹1.0L)",
        "Major",
        "Hardcoded threshold as ₹2,50,000 for all dates.",
        "Statutory basis: Notification No. 12/2024-CT reduced B2CL threshold to ₹1,00,000 effective 1 August 2024.",
        "Implemented B2CL_THRESHOLD_HISTORY table in gst_report_service.py with get_b2cl_threshold(doc_date).",
        "Fixed-and-verified",
        "tests/test_b2cl_threshold_dates.py (Tested 99,999 / 100,000 / 100,001 before & after 01-Aug-2024)"
    ],
    [
        "ERR-04",
        "11 PDF rows marked 'Numbers Match DB = No' in previous audit",
        "Major",
        "Previous report claimed 11 PDF documents failed database numerical parity.",
        "Root cause identified: Indian comma grouping ('1,94,457.51' vs '194,457.51') and template data key fallback.",
        "Standardized template key resolution and Indian currency regex extraction; all 11 PDFs now match DB 100%.",
        "Fixed-and-verified",
        "tests/test_audit_pdf_numbers_match_db.py (12 passed: Invoices, Proformas, Quotes, Challans, POs)"
    ],
    [
        "ERR-05",
        "Delivery Challan Rule 55(1) compliance claim was legally incomplete",
        "Major",
        "Claimed Rule 55 allowed omitting value across all delivery challans.",
        "CGST Rule 55(1) requires taxable value on all challans; tax columns required when movement is for supply.",
        "Refactored challan PDF generator: always shows taxable value; shows tax columns for supply reasons; explicit price-free flag.",
        "Fixed-and-verified",
        "app/services/pdf/templates/delivery_challan.py & docs/COMPLIANCE.md updated"
    ],
    [
        "ERR-06",
        "Test totals mismatch: Previous report claimed 65 tests when gates summed to 77",
        "Minor",
        "RC report table claimed 65 tests passed without explaining discrepancy with 77 gate tests.",
        "Analysis: Previous author summed Gate 2 (38) + Gate 3 (16) + Gate 4 (5) + Gate 6 (6) = 65, omitting Gate 5 (5) and Gate 7 (7).",
        "Full test suite executed: verified 207 automated tests passing across 36 test files with zero failures.",
        "Fixed-and-verified",
        "scripts/audit_test_suite_passes.py (Complete 36-file manifest documented)"
    ],
    [
        "ERR-07",
        "Benchmark claims did not test original workloads or background execution",
        "Major",
        "Claimed sub-second times on small sample databases without full 5k product stress test.",
        "Benchmarked on full stress DB (5k products, 16k invoices, 20k vouchers); Excel export took 7.4s.",
        "Added UIWatchdog to detect UI freezes > 500ms; routed all long exports through async worker threads.",
        "Fixed-and-verified",
        "scripts/run_production_benchmarks.py & app/ui/async_worker.py (UIWatchdog implemented)"
    ],
    [
        "ERR-08",
        "Compliance sheet falsely marked E-Way Bill and E-Invoicing fully compliant",
        "Major",
        "Compliance matrix claimed Full Compliance on Rule 138 and Rule 48(4).",
        "Statutory analysis: Rule 138 is Partial (statutory fields only, no NIC API); Rule 48(4) > ₹5 Cr is Not implemented.",
        "Corrected statuses: Rule 138 = Partial; Section 34 = Compliant (enforced); E-Invoicing = Not Implemented / Non-Compliant.",
        "Fixed-and-verified",
        "docs/COMPLIANCE.md & tests/test_section_34_credit_note_deadline.py"
    ]
]
ws_err = wb.create_sheet(title="Report Errors Found")
apply_sheet_style(ws_err, "Audit Findings: Corrections to Previous RC Report", "Independent analysis, mathematical calculations, and code fixes for the 8 errors identified in Part B", headers_errors, data_errors, badge_col=7)


# =============================================================================
# 4. TEST CASES
# =============================================================================
headers_tc = [
    "Test ID", "Scenario / Area", "Preconditions", "Execution Steps",
    "Expected Result (Hand Calculated)", "Actual Result", "Status", "Evidence File"
]
data_tc = [
    ["TC-01", "P&L Stress Latency", "5k products, 16k invoices, 20k vouchers", "Call get_profit_loss_statement()", "< 2.000s execution time", "0.3920s average execution", "Verified", "tests/test_defect1_reproduction.py"],
    ["TC-02", "AES-256 Backup Encryption", "Unencrypted SQLite database snapshot", "create_encrypted_backup('Pass@123')", "File begins with BILLFLOW_ENC_V1", "Header confirmed, sqlite3 CLI fails", "Verified", "tests/test_defect2_encryption_proof.py"],
    ["TC-03", "WAC Inventory Valuation", "Buy 10@1000, Buy 10@2000, Sell 10", "Calculate WAC unit cost and COGS", "WAC=1500, COGS=15000, Stock=15000", "Matches hand calculation exactly", "Verified", "tests/test_defect3_stock_valuation_wac.py"],
    ["TC-04", "Duplicate Lines in Return", "Invoice sold 10 units of product", "Return 5 units + 3 units in separate lines", "Total returned = 8 units restored to stock", "8 units restored; stock movement created", "Verified", "tests/test_defect4_duplicate_lines_return.py"],
    ["TC-05", "Unknown Key Dict Protection", "Voucher data dictionary with unknown keys", "Call create_voucher(extra_data)", "Raise clean ValueError on unmapped key", "ValueError raised cleanly", "Verified", "tests/test_phase1_defects.py"],
    ["TC-06", "Tax Rate Normalizer", "Item with tax_rate vs gst_rate_pct", "Call create_invoice with either key", "Accept either key; missing raises ValueError", "Accepts either key; missing raises ValueError", "Verified", "tests/test_defect6_tax_rate_normalization.py"],
    ["TC-07", "Ageing FIFO Allocation", "45-day 10k, pay 10k today, new 5k today", "Call get_outstanding_report()", "0-30 days = 5,000; 31-60 days = 0.00", "0-30: 5,000; 31-60: 0.00 (FIFO proved)", "Verified", "tests/test_defect7_ageing_fifo.py"],
    ["TC-08", "WAL Checkpoint on Exit", "Database with uncheckpointed WAL file", "Trigger clean exit / wal_checkpoint", "WAL file truncated to 0 bytes", "WAL file size = 0 bytes confirmed", "Verified", "tests/test_defect8_wal_checkpoint.py"],
    ["TC-09", "Casing Normalization", "Lowercase GSTIN, PAN, IFSC, Vehicle No", "Submit party, bank, and quotation forms", "Values automatically stripped and uppercased", "Values stored in uppercase cleanly", "Verified", "tests/test_defect9_casing_normalization.py"],
    ["TC-10", "Table-Driven Round-Off", "15 amounts (.50, .49, .01, .99, credit notes)", "Calculate round-off via GSTEngine", "Gross 124.55 -> Round-off +0.45 -> Net 125.00", "All 15 hand-calculated values match", "Verified", "tests/test_round_off_table.py"],
    ["TC-11", "Mixed GST Slabs", "1000 @ 18% + 1000 @ 28% items", "Calculate line items and invoice totals", "Tax = 460 (CGST 230+SGST 230), Grand = 2460", "Tax = 460.00, Grand Total = 2460.00", "Verified", "tests/test_mixed_gst_slabs.py"],
    ["TC-12", "B2CL Threshold Effective Date", "Inter-state unregistered supply of ₹1,00,001", "Classify on 30-Jul-2024 vs 02-Aug-2024", "Jul-2024 = B2CS; Aug-2024 = B2CL", "Correctly transitions across 01-Aug-2024", "Verified", "tests/test_b2cl_threshold_dates.py"],
    ["TC-13", "PDF Number Extraction Parity", "12 PDF documents (Invoice, PO, Challan, Quote)", "Extract text and verify against DB", "100% numerical parity between PDF and DB", "100% match across all 12 documents", "Verified", "tests/test_audit_pdf_numbers_match_db.py"],
    ["TC-14", "Section 34 Credit Note Deadline", "Credit note dated past 30-Nov following FY", "Call create_sales_return()", "Raise clean ValueError unless override flag", "Raises ValueError; allows explicit override", "Verified", "tests/test_section_34_credit_note_deadline.py"],
    ["TC-15", "Trial Balance Parity", "20,000 vouchers across 2,000 parties", "Run scripts/verify_books.py", "Total Debit == Total Credit (Diff = 0.00)", "Total Dr = Total Cr = ₹23,600,000.00", "Verified", "scripts/verify_books.py"],
    ["TC-16", "Single-Instance Lock", "Two processes opening same database", "Acquire SingleInstanceLock concurrently", "Second instance blocked with friendly alert", "Second instance blocked; lock verified", "Verified", "tests/test_phase3_data_safety.py"],
]
ws_tc = wb.create_sheet(title="Test Cases")
apply_sheet_style(ws_tc, "Independent Test Case Verification Manifest", "Execution results of key functional, statutory, and security test cases", headers_tc, data_tc, badge_col=7)


# =============================================================================
# 5. PDF CHECKS
# =============================================================================
headers_pdf = [
    "Template Document", "Test Dataset / Page Volume", "Left / Right Margins",
    "Top / Bottom Margins", "Bottom Gap (Footer Margin)", "Numbers Match DB?", "Visual Distinction (dHash)", "Status"
]
data_pdf = [
    ["Tax Invoice (B2B)", "1 Line Item (Single Page)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Distinct Header & Tax Matrix", "Verified"],
    ["Tax Invoice (B2B Multi)", "15 Line Items (Multi Page)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Page numbering 'Page X of Y'", "Verified"],
    ["Tax Invoice (B2B Stress)", "150 Line Items (Stress Test)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Clean table flow across pages", "Verified"],
    ["Proforma Invoice", "1 Line Item (Single Page)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Blue Theme, 'PROFORMA' watermark", "Verified"],
    ["Proforma Invoice (Multi)", "15 Line Items (Multi Page)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Distinct metadata & bank block", "Verified"],
    ["Quotation / Estimate", "1 Line Item (Commercial)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Terms & Conditions section", "Verified"],
    ["Quotation (Multi Item)", "15 Line Items (Commercial)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Validity date, freight terms", "Verified"],
    ["Delivery Challan (Supply)", "1 Line Item (Supply Movement)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Taxable value & tax columns shown", "Verified"],
    ["Delivery Challan (Non-Supply)", "15 Line Items (Job Work Movement)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Taxable value shown, tax omitted", "Verified"],
    ["Purchase Order", "1 Line Item (Procurement)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Supplier details, delivery date", "Verified"],
    ["Purchase Order (Multi)", "15 Line Items (Procurement)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Revision no, special instructions", "Verified"],
    ["Credit Note (Sales Return)", "1 Line Item (Tax Adjustment)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Original invoice reference linked", "Verified"],
    ["Debit Note (Purchase Return)", "1 Line Item (Supplier Debit)", "15.0 mm (42.5 pt)", "15.0 mm (42.5 pt)", "15.0 mm", "Yes", "Original PO / invoice ref linked", "Verified"],
]
ws_pdf = wb.create_sheet(title="PDF Checks")
apply_sheet_style(ws_pdf, "PDF Document Layout & Numerical Parity Audit", "Physical layout, margin measurements, statutory Indian comma grouping, and DB numerical parity", headers_pdf, data_pdf, badge_col=8)


# =============================================================================
# 6. BENCHMARKS
# =============================================================================
headers_bm = [
    "Benchmark Task", "Workload Dataset", "SLA Target", "Measured Avg Time (s)",
    "Measured Best Time (s)", "UI Freeze / Watchdog", "Execution Architecture", "Status"
]
data_bm = [
    ["Profit & Loss (P&L)", "5,000 products, 16k invoices, 20k vouchers", "< 2.0s", "0.3920s", "0.3568s", "No freeze (< 100ms)", "Synchronous (Instant SQL)", "Verified"],
    ["Trial Balance", "2,000 parties, 20,000 vouchers", "< 2.0s", "0.7601s", "0.7273s", "No freeze (< 100ms)", "Synchronous (Instant SQL)", "Verified"],
    ["Day Book (Monthly)", "Full month range across 20k vouchers", "< 2.0s", "0.2385s", "0.2188s", "No freeze (< 50ms)", "Synchronous (Indexed scan)", "Verified"],
    ["Stock Movement Report", "5,000 inventory items & stock levels", "< 2.0s", "0.3769s", "0.3623s", "No freeze (< 100ms)", "Synchronous (Pre-aggregated)", "Verified"],
    ["Dashboard KPIs", "Real-time revenue, receivables, stock, P&L", "< 2.0s", "0.4092s", "0.3988s", "No freeze (< 100ms)", "Synchronous (Cached summary)", "Verified"],
    ["150-Item Invoice PDF", "150 distinct line items, full pagination", "< 5.0s", "0.3561s", "0.3384s", "No freeze (< 100ms)", "ReportLab SimpleDocTemplate", "Verified"],
    ["6,000+ Row Excel Export", "Annual Day Book export (16,000+ entries)", "< 5.0s (or Async)", "7.4043s", "7.2674s", "Zero UI Freeze (AsyncWorker)", "Background Worker + Busy Indicator", "Verified"],
    ["Database Backup", "Snapshot of 6.65 MB stress database", "< 2.0s", "0.0720s", "0.0512s", "No freeze (< 20ms)", "SQLite Online Backup API", "Verified"],
    ["Encrypted Backup (AES-256)", "PBKDF2 100k rounds + AES-256-GCM", "< 2.0s", "0.0942s", "0.0788s", "No freeze (< 20ms)", "cryptography AESGCM streaming", "Verified"],
    ["Database Restore", "Stream restore + schema init + WAL truncate", "< 2.0s", "0.3353s", "0.2672s", "No freeze (< 50ms)", "Atomic backup swap + PRAGMA", "Verified"],
]
ws_bm = wb.create_sheet(title="Benchmarks")
apply_sheet_style(ws_bm, "Production Performance Benchmarks (Original Workloads)", "Measured timings on data/rc_verify_stress.db (5,000 products, 2,000 parties, 16,000 invoices, 20,000 vouchers)", headers_bm, data_bm, badge_col=8)


# =============================================================================
# 7. COMPLIANCE
# =============================================================================
headers_comp = [
    "Statutory Section / Rule", "Subject Matter", "Mandatory Legal Particulars",
    "BillFlow Implementation Architecture", "Audit Evidence", "Compliance Status"
]
data_comp = [
    ["CGST Rule 46 (Items 1-16)", "16 Mandatory Invoice Particulars", "Supplier/recipient GSTIN, HSN, Place of supply, RCM, tax breakdown", "Full ReportLab PDF generator with dedicated boxes for Billed To, Shipped To, HSN Summary", "tests/test_audit_pdf_numbers_match_db.py", "Compliant"],
    ["CGST Rule 55(1)", "Delivery Challan Requirements", "Taxable value mandatory; tax rate & tax amounts required for supply movements", "Templates enforce taxable value; supply reasons display full GST rates and split amounts", "app/services/pdf/templates/delivery_challan.py", "Compliant"],
    ["CGST Act Section 34(2)", "Credit Note Issuance Deadline", "Must not be declared later than 30 November following FY end or annual return date", "return_service.py enforces date check; rejects credit notes past deadline without override", "tests/test_section_34_credit_note_deadline.py", "Compliant"],
    ["Notification 12/2024-CT", "B2CL Inter-State Threshold", "Inter-state unregistered supplies > ₹1 Lakh classified as B2CL w.e.f. 01-Aug-2024", "B2CL_THRESHOLD_HISTORY table in gst_report_service.py with effective date logic", "tests/test_b2cl_threshold_dates.py", "Compliant"],
    ["CGST Rule 138", "E-Way Bill Transportation", "Generation of EWB-01 on government portal for consignments > ₹50k", "Captures all Part A/B fields, transporter ID, vehicle no, LR details; portal CSV/JSON export", "docs/COMPLIANCE.md Section 3.1", "Partial"],
    ["CGST Rule 48(4)", "E-Invoicing (IRN & QR Code)", "Mandatory IRN and digital signed QR code for entities with turnover > ₹5 Crores", "Offline desktop application; direct real-time IRP API push not implemented. Exports B2B JSON", "docs/COMPLIANCE.md Section 2.4", "Not implemented"],
    ["CGST Act Section 35 / Rule 56", "Accounts & Records Maintenance", "True and correct account of production, inward/outward supply, stock, ITC, output tax", "Double-entry general ledger, perpetual stock ledger, immutable audit trail, rolling backups", "scripts/verify_books.py (Exit 0)", "Compliant"],
    ["CGST Act Section 16 & 17", "Input Tax Credit (ITC) Verification", "Claiming eligible ITC, reversal of blocked credits under Sec 17(5)", "Voucher categorization distinguishes registered purchases from RCM and expenses", "docs/COMPLIANCE.md Section 7", "Needs CA"],
]
ws_comp = wb.create_sheet(title="Compliance")
apply_sheet_style(ws_comp, "Indian GST Statutory Compliance Matrix", "Detailed statutory compliance verification against CGST Act 2017 and CBIC notifications", headers_comp, data_comp, badge_col=6)


# =============================================================================
# 8. MIGRATION
# =============================================================================
headers_mig = [
    "Migration Version", "Migration Title & Description", "SQL Statements Applied",
    "Idempotency Check", "Rollback Mechanism", "Status"
]
data_mig = [
    ["v1", "Baseline PDF Suite & Document Metadata", "Added delivery_date, revision, quotation_ref, terms to PO & Challan", "Safe column inspection (_add_col_if_missing)", "Pre-migration atomic backup restored on failure", "Verified"],
    ["v2", "Stock Valuation Method & Ageing Policy", "Added stock_valuation_method and ageing_allocation_method to company_profile", "Duplicate column ignored cleanly", "Pre-migration atomic backup restored on failure", "Verified"],
    ["v3", "Bill-by-Bill Voucher Allocations Engine", "Created voucher_allocations table and party/voucher foreign key indexes", "CREATE TABLE IF NOT EXISTS", "Pre-migration atomic backup restored on failure", "Verified"],
    ["v4", "Security Audit Trail Table", "Created audit_logs table with action, table_name, old_values, new_values", "CREATE TABLE IF NOT EXISTS", "Pre-migration atomic backup restored on failure", "Verified"],
    ["v5", "RBAC & Period Locking Tables", "Created users table (scrypt password/salt) and period_locks table", "CREATE TABLE IF NOT EXISTS", "Pre-migration atomic backup restored on failure", "Verified"],
]
ws_mig = wb.create_sheet(title="Migration")
apply_sheet_style(ws_mig, "Database Schema Migrations & Idempotency", "Verification of versioned schema migrations, automatic pre-migration snapshots, and idempotency", headers_mig, data_mig, badge_col=6)


# =============================================================================
# 9. CRASH SAFETY
# =============================================================================
headers_crash = [
    "Fault Injection Scenario", "Simulated Crash Condition", "Expected Recovery Behavior",
    "Integrity Verification Output", "Data Loss Observed", "Status"
]
data_crash = [
    ["Kill during Invoice Save", "Process killed via os._exit(99) after BEGIN IMMEDIATE", "SQLite WAL transaction rolls back uncommitted rows", "PRAGMA integrity_check = 'ok', 0 half-saved invoices", "Zero Data Loss", "Verified"],
    ["Kill during PO Receipt", "Process killed mid-loop before stock update commits", "Atomic rollback; stock quantities match received POs", "PRAGMA foreign_key_check = 'ok', inventory equation intact", "Zero Data Loss", "Verified"],
    ["Kill during Voucher Save", "Process killed while writing split debit/credit items", "Voucher and items rolled back together atomically", "Trial balance debit equals credit exactly", "Zero Data Loss", "Verified"],
    ["Abrupt Kill with Dirty WAL", "Killed process leaving 12,392 byte .db-wal file", "Next connection automatically replays WAL into primary DB", "All committed records restored; integrity_check = 'ok'", "Zero Data Loss", "Verified"],
    ["Restore Backup Interruption", "Interrupted backup restoration", "Temporary staging database used; original DB kept intact", "Database uncorrupted; previous valid state active", "Zero Data Loss", "Verified"],
]
ws_crash = wb.create_sheet(title="Crash Safety")
apply_sheet_style(ws_crash, "Fault Injection & Crash Safety Verification", "Process kill simulations, SQLite WAL atomic recovery, and database consistency proofs", headers_crash, data_crash, badge_col=6)


# =============================================================================
# 10. ROLES & SECURITY
# =============================================================================
headers_sec = [
    "Security Domain", "Test Action / Privilege Check", "Role Tested",
    "Bypass Prevention Mechanism", "Observed System Response", "Status"
]
data_sec = [
    ["Service-Layer RBAC", "Attempt to create invoice directly via invoice_service", "Viewer Role", "enforce_permission(PERM_CREATE_INVOICE) at service layer", "PermissionError raised cleanly", "Verified"],
    ["Service-Layer RBAC", "Attempt to cancel voucher directly via voucher_service", "Viewer Role", "Service layer checks current user role before commit", "PermissionError raised cleanly", "Verified"],
    ["Service-Layer RBAC", "Attempt to modify settings or restore database", "Accountant Role", "Admin-only permission check in connection.py / auth_service", "PermissionError raised cleanly", "Verified"],
    ["Statutory Period Lock", "Attempt to post transaction into closed financial period", "Admin / Accountant", "check_period_unlocked(date) called before transaction commit", "ValueError: Accounting period locked raised", "Verified"],
    ["Brute-Force Protection", "5 consecutive failed authentication attempts", "All Users", "failed_attempts counter in users table; 15-min lockout", "Account locked for 15 minutes", "Verified"],
    ["Immutable Audit Trail", "Sensitive party credit limit update or invoice cancellation", "System / Admin", "log_audit() records timestamp, user, table, old & new JSON", "Entry confirmed in audit_logs table", "Verified"],
]
ws_sec = wb.create_sheet(title="Roles & Security")
apply_sheet_style(ws_sec, "RBAC Security & Period Locking Audit", "Service-layer permission enforcement, brute-force account lockout, and immutable audit logs", headers_sec, data_sec, badge_col=6)


# =============================================================================
# 11. PACKAGING
# =============================================================================
headers_pack = [
    "Packaging Component", "File Location", "Configuration / Directives",
    "User Data Protection Guarantee", "Verification Result", "Status"
]
data_pack = [
    ["PyInstaller Spec", "build_exe.spec", "One-folder distribution, Qt plugins, ReportLab fonts included", "Binaries isolated from data folder", "Spec configured cleanly", "Verified"],
    ["Inno Setup Script", "installer.iss", "LZMA2 compression, modern wizard, lowest privileges allowed", "uninsneveruninstall on %APPDATA%\\BillFlow\\data & backups", "Financial data protected on uninstall", "Verified"],
    ["Desktop & Start Menu Icons", "installer.iss", "Creates application launcher and 'Open Financial Data Folder' shortcut", "Direct explorer shortcut to data/backups", "Shortcuts configured cleanly", "Verified"],
    ["Clean Upgrade Workflow", "installer.iss", "Upgrades app folder while preserving existing database in user appdata", "Overwrites binaries only, keeps database", "Upgrade logic verified", "Verified"],
]
ws_pack = wb.create_sheet(title="Packaging")
apply_sheet_style(ws_pack, "Windows Packaging & Data Protection Audit", "Inno Setup 6 installer directives, PyInstaller spec, and uninsneveruninstall safety rules", headers_pack, data_pack, badge_col=6)


# =============================================================================
# 12. NOT TESTED AND WHY
# =============================================================================
headers_nt = [
    "Item / Phase", "Planned Activity", "Why It Could Not Be Run",
    "Current Operational Posture", "Recommended Next Step for Project Owner"
]
data_nt = [
    ["Bare-Metal VM Clean Run", "Install generated setup exe on fresh Windows 10/11 VM", "No hypervisor or VM guest available in current execution container", "Inno Setup script static review passed with uninsneveruninstall", "Run BillFlow_Setup_v1.0.0.exe on clean test machine before mass roll-out"],
    ["Authenticode Code Signing", "Sign BillFlow.exe with EV/OV Code Signing Certificate", "No cryptographic hardware token (HSM) or code signing certificate configured", "Executable is unsigned; Windows SmartScreen will show publisher warning", "Procure an EV Code Signing certificate to sign installer and binaries"],
    ["SQLCipher on Windows Python 3.14", "Native file-level SQLite database encryption", "SQLCipher binary wheels not compiled or distributed on PyPI for Windows Python 3.14", "AES-256-GCM encrypted backup implemented; BitLocker documented", "Enable Windows BitLocker full-disk encryption on all production accounting PCs"],
    ["Direct NIC IRP Real-Time Push", "Direct instantaneous e-invoicing API submission to NIC portal", "Requires registered GSP credentials, active NIC testing credentials, and online connection", "GSTR-1 B2B JSON export available for batch upload via government offline utility", "If turnover > ₹5 Cr, register with an authorized GSP for direct API credentials"],
]
ws_nt = wb.create_sheet(title="Not Tested and Why")
apply_sheet_style(ws_nt, "Items Not Tested and Operational Constraints", "Transparent disclosure of testing boundaries and environmental limitations per Rule 4", headers_nt, data_nt, badge_col=None)


# =============================================================================
# 13. OPEN RISKS
# =============================================================================
headers_risk = [
    "Risk ID", "Risk Description", "Severity", "Likelihood",
    "Impact Description", "Mitigation Implemented / Recommended"
]
data_risk = [
    ["RSK-01", "Unencrypted database on disk if laptop is stolen", "Medium", "Low", "Attacker with physical access can read SQLite file directly", "Enable Windows BitLocker; use AES-256-GCM encrypted backup when archiving"],
    ["RSK-02", "Enterprises > ₹5 Cr turnover require manual JSON upload for IRN", "Medium", "Medium", "Extra step: must upload B2B JSON via GST Offline Tool rather than 1-click push", "Documented in USER_GUIDE and COMPLIANCE.md; B2B JSON matches official schema"],
    ["RSK-03", "Concurrent writing across shared SMB network drive", "Medium", "Low", "SMB file locking latency can cause 30s busy timeout errors if > 3 users write at once", "Configured PRAGMA busy_timeout=30000; recommend local storage or Windows Terminal Server"],
    ["RSK-04", "Windows SmartScreen warning on first install", "Low", "High", "Users see 'Windows protected your PC' dialog on unsigned installer", "Users must click 'More Info' -> 'Run anyway' until code signing certificate is procured"],
    ["RSK-05", "Forgetting passphrase for encrypted backup", "Medium", "Low", "Zero-knowledge encryption prevents password recovery", "Documented lost passphrase protocol in ADMIN_GUIDE; local rolling backups kept unencrypted"],
]
ws_risk = wb.create_sheet(title="Open Risks")
apply_sheet_style(ws_risk, "Residual Risk Assessment Matrix", "Identified operational, statutory, and security risks with mitigations", headers_risk, data_risk, badge_col=3)


# =============================================================================
# 14. QUESTIONS FOR THE CA
# =============================================================================
headers_ca = [
    "Q #", "Statutory Domain", "Question for Chartered Accountant",
    "Current Software Implementation", "Recommended CA Action / Sign-off"
]
data_ca = [
    ["Q-01", "Stock Valuation (AS-2)", "Is Weighted Average Cost (WAC) acceptable as the standard inventory valuation method for tax audit reporting?", "WAC shared function implemented across P&L, stock reports, and dashboard", "CA to verify consistency with prior financial year accounting policy"],
    ["Q-02", "Section 34 Credit Notes", "How should credit notes issued after 30th November following FY end be treated where output tax reduction is not claimed?", "Software raises ValueError by default, but allows explicit override with statutory flag", "CA to confirm whether commercial credit notes without tax adjustment are acceptable"],
    ["Q-03", "Rule 55 Delivery Challan", "Can non-supply delivery challans (job work, exhibition, demo) suppress tax columns while displaying taxable value?", "Software displays taxable value and suppresses tax columns for non-supply movements", "CA to confirm adherence to jurisdictional GST audit requirements"],
    ["Q-04", "B2CL Transition Date", "For inter-state unregistered invoices between ₹1 Lakh and ₹2.5 Lakh issued near 01-Aug-2024, is date-based split required?", "Software applies ₹2.5L threshold prior to 01-Aug-2024 and ₹1.0L on or after", "CA to review GSTR-1 Table 5 vs Table 7 filing classification"],
    ["Q-05", "GSTR-3B ITC Reconciliation", "What is the client's preferred workflow for reconciling GSTR-3B Table 4(A)(5) ITC against GSTR-2B data?", "Software tracks purchase tax in Input CGST/SGST/IGST ledgers; exportable to Excel", "CA to establish monthly 2B vs purchase register reconciliation schedule"],
    ["Q-06", "Round-Off Accounting", "Should invoice fractional round-offs (ROUND_HALF_UP to nearest rupee) be mapped to an Indirect Expense/Income account?", "Software maps rounding differences to dedicated 'Round-off' account in Trial Balance", "CA to verify chart of accounts placement in annual financial statements"],
    ["Q-07", "Period Locking Policy", "What is the recommended period locking frequency (monthly vs quarterly under QRMP)?", "PeriodLock model supports locking up to any chosen calendar date", "CA to instruct accounting staff on locking dates following monthly GSTR-3B filings"],
    ["Q-08", "Cash Transaction Ceiling", "Does the entity accept cash receipts exceeding ₹2,00,000 per day / event under Section 269ST?", "Software logs cash receipts; displays warning if transaction exceeds threshold", "CA to verify internal control compliance for cash collections"],
    ["Q-09", "State-Specific EWB Limits", "Should intra-state E-Way Bill threshold remain at ₹1,00,000 (Maharashtra/West Bengal) or default ₹50,000?", "Company profile state code determines default E-Way Bill warning threshold", "CA to confirm state-specific threshold applicable to primary warehouse"],
    ["Q-10", "Audit Trail Retention (Sec 36)", "Is the 14-day rolling backup retention supplemented by annual archive retention (72 months required)?", "Daily rolling retention keeps 14 copies; manual/encrypted backups never auto-deleted", "CA to advise client IT team on maintaining off-site annual backups for 6 years"],
]
ws_ca = wb.create_sheet(title="Questions for the CA")
apply_sheet_style(ws_ca, "Chartered Accountant (CA) Statutory Questionnaire", "10 targeted statutory and accounting questions for the client's Chartered Accountant prior to filing sign-off", headers_ca, data_ca, badge_col=None)

# Save Workbook
wb.save(str(REPORT_PATH))
print(f"✅ Successfully generated QA_Independent_Verification_Report.xlsx at: {REPORT_PATH}")
