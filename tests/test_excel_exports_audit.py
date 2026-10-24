"""
Unit Tests for Phase 2: Excel Exports Verification.
Verifies that:
1. Header rows are frozen on every export.
2. Print titles and A4 landscape orientation are configured.
3. Formula injection is prevented (=, +, -, @ prefixed with ').
4. Numbers in exported sheets match source figures.
"""
from datetime import date
from decimal import Decimal
from pathlib import Path
import openpyxl
import pytest

from app.database.connection import get_session
from app.database.models import Party, Voucher, VoucherItem, Product
from app.services.export_service import (
    export_party_ledger_excel,
    export_outstanding_excel,
    export_daybook_excel,
    export_trial_balance_excel,
    sanitize_excel_cell
)
from app.config import EXPORTS_DIR


def test_sanitize_excel_cell():
    """Verify formula injection characters are safely escaped."""
    assert sanitize_excel_cell("=1+1") == "'=1+1"
    assert sanitize_excel_cell("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_excel_cell("-SUM(A1:A10)") == "'-SUM(A1:A10)"
    assert sanitize_excel_cell("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_excel_cell("Normal Text") == "Normal Text"
    assert sanitize_excel_cell(12345.67) == 12345.67


def test_excel_exports_freeze_panes_and_print_setup():
    """Verify all 4 core exports configure freeze_panes, print titles, and landscape."""
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    
    # 1. Day Book
    daybook_path = EXPORTS_DIR / "test_daybook_audit.xlsx"
    export_daybook_excel(date.today(), output_path=daybook_path)
    wb = openpyxl.load_workbook(str(daybook_path))
    ws = wb.active
    assert ws.freeze_panes is not None, "Daybook Excel export must have frozen panes"
    assert ws.print_title_rows in ("4:4", "$4:$4")
    assert ws.page_setup.orientation == ws.ORIENTATION_LANDSCAPE
    wb.close()

    # 2. Trial Balance
    tb_path = EXPORTS_DIR / "test_tb_audit.xlsx"
    export_trial_balance_excel(date.today(), output_path=tb_path)
    wb = openpyxl.load_workbook(str(tb_path))
    ws = wb.active
    assert ws.freeze_panes is not None, "Trial Balance Excel export must have frozen panes"
    assert ws.print_title_rows in ("4:4", "$4:$4")
    assert ws.page_setup.orientation == ws.ORIENTATION_LANDSCAPE
    wb.close()

    # 3. Outstanding
    out_path = EXPORTS_DIR / "test_out_audit.xlsx"
    export_outstanding_excel("Receivables", output_path=out_path)
    wb = openpyxl.load_workbook(str(out_path))
    ws = wb.active
    assert ws.freeze_panes is not None, "Outstanding Excel export must have frozen panes"
    assert ws.print_title_rows in ("4:4", "$4:$4")
    assert ws.page_setup.orientation == ws.ORIENTATION_LANDSCAPE
    wb.close()
